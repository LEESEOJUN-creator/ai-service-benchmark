"""Private, dated PDF/XLSX/chart export from the append-only execution journal."""
from datetime import date
from pathlib import Path
import os
import re
import tempfile
import zipfile
import math
import unicodedata
from xml.sax.saxutils import escape

from .costs import converted_cost, krw_text, normalize_rates
from .session import read_session
from .overview import report_groups, coverage_text, summaries, assessment_warnings

SERVICE_NAMES = {"monogpt": "MonoGPT", "timely": "Timely AI", "chatgpt": "ChatGPT"}


def _text(value):
    return "미확인" if value is None or value == "" else str(value)


def _time(result):
    low, high = result.get("last_generating_observed_ms"), result.get("first_complete_observed_ms")
    if high is None:
        return "미관측"
    return (f"≤{high / 1000:.1f}초 (완료 관측 상한)" if low is None
            else f"{low / 1000:.1f}~{high / 1000:.1f}초 (관측 구간)")


def _cost(result, service, rates):
    if result.get("cost") is None:
        return "미노출" if result.get("cost_status") == "unavailable" else "미확정"
    return krw_text(converted_cost(result, service, rates)["krw"])


def _rows(doc):
    results = {e["request_id"]: e for e in doc["events"] if e["kind"] == "result"}
    assessments = {e["request_id"]: e for e in doc["events"] if e["kind"] == "assessment"}
    cases = {c["id"]: c for c in doc["cases"]}
    rows = []
    for request in (e for e in doc["events"] if e["kind"] == "reserve"):
        key = request["request_id"]
        result, assessment = results.get(key, {}), assessments.get(key)
        if result.get("status") == "success" and assessment is None:
            raise ValueError(f"{key}: 기준별 관찰·판단·근거가 없어 최종 보고서를 만들 수 없음")
        rows.append({"request": request, "result": result, "assessment": assessment,
                     "case": cases[request["case_id"]], "group": request.get("model_group") or request["model"]})
    return rows


def _literal(cell, value):
    # Copied answers beginning with '=' must never become spreadsheet formulas.
    cell.value = value[:32767] if isinstance(value, str) else value
    if isinstance(value, str):
        cell.data_type = "s"


def _cache_ledger_totals(path, totals, costs, ledger_index=3):
    """Cache evaluated score/cost formulas for previewers before Excel opens."""
    cached = path.with_suffix(".cached.xlsx")
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(cached, "w") as target:
        for info in source.infolist():
            payload = source.read(info.filename)
            if info.filename == f"xl/worksheets/sheet{ledger_index}.xml":
                for column, entries in (("I", totals), ("M", costs)):
                    for row_no, total in entries.items():
                        pattern = (rb'(<c r="' + column.encode() + str(row_no).encode() +
                                   rb'"[^>]*><f>.*?</f>)(?:<v\s*/>|<v></v>)')
                        payload, count = re.subn(pattern, lambda match: match[1] +
                                                 b"<v>" + str(total).encode() + b"</v>",
                                                 payload, count=1, flags=re.DOTALL)
                        if count != 1:
                            raise ValueError(f"수식 캐시 위치를 찾을 수 없음: {column}{row_no}")
            target.writestr(info, payload)
    os.replace(cached, path)


def _write_xlsx(path, rows, doc, today, name, rates):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    compare = wb.active
    compare.title = "서비스 비교"
    detail = wb.create_sheet("기준별 평가")
    ledger = wb.create_sheet("실측 원장")
    answers = wb.create_sheet("답변 원문")
    ux = wb.create_sheet("사용 경험")
    navy, stripe = "193B63", "F5F8FC"

    def setup(ws, labels, widths):
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(labels))
        ws.cell(1, 1, ws.title).font = Font(name="Arial", size=17, bold=True, color=navy)
        ws.row_dimensions[1].height = 32
        ws.cell(2, 1, f"생성일 {today.isoformat()} · 실행 {name} · UI 표시 조건 비교, 백엔드 동일성 미검증")
        ws.cell(2, 1).font = Font(name="Arial", size=10, color="52657A")
        ws.sheet_view.showGridLines = False
        ws.freeze_panes = "D5"
        for col, (label, width) in enumerate(zip(labels, widths), 1):
            cell = ws.cell(4, col, label)
            cell.fill = PatternFill("solid", fgColor=navy)
            cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(col)].width = width
        ws.row_dimensions[4].height = 32

    def add(ws, values, height=52):
        row_no = ws.max_row + 1
        lines = 1
        for col, value in enumerate(values, 1):
            cell = ws.cell(row_no, col)
            _literal(cell, value)
            cell.font = Font(name="Arial", size=10, color="24364B")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=Side(style="thin", color="D9E2EC"))
            if row_no % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=stripe)
            width = max(10, ws.column_dimensions[get_column_letter(col)].width - 3)
            wrapped = sum(max(1, math.ceil(sum(2 if unicodedata.east_asian_width(ch) in 'WF' else 1
                                              for ch in part) / width)) for part in str(value or '').split('\n'))
            lines = max(lines, wrapped)
        ws.row_dimensions[row_no].height = min(409, max(32, lines * 14 + 12))
        return row_no

    services = list(SERVICE_NAMES) + list(dict.fromkeys(
        r["request"]["service"] for r in rows if r["request"]["service"] not in SERVICE_NAMES))
    setup(compare, ["모델 묶음", "질문 ID", *[SERVICE_NAMES.get(s, s) for s in services], "서비스 간 차이·실무 영향"],
          [26, 23, *([65] * len(services)), 85])
    comparisons = {(e['model_group'], e['case_id']): e for e in doc['events'] if e['kind'] == 'comparison'}
    grouped = {}
    for row in rows:
        grouped.setdefault((row["group"], row["case"]["id"]), {})[row["request"]["service"]] = row
    planned_cases = {g['label']: set(g['cases']) for g in report_groups(rows, doc)}
    for group in dict.fromkeys(row["group"] for row in rows):
        for case in doc["cases"]:
            if group in planned_cases and case['id'] not in planned_cases[group]:
                continue
            grouped.setdefault((group, case["id"]), {})
    group_members = {}
    for row in rows:
        group_members.setdefault(row["group"], []).append(row)
    for (group, case_id), members in grouped.items():
        values = [group, case_id]
        for service in services:
            item = members.get(service)
            if item is None:
                exclusion = next((r["request"].get("excluded_services", {}).get(service)
                                  for r in group_members[group]
                                  if r["request"].get("excluded_services", {}).get(service)), None)
                values.append("모델 없음 — " + exclusion if exclusion else
                              "기록 없음 — 모델 없음으로 추정하지 않음")
                continue
            req, res, assessment = item["request"], item["result"], item["assessment"]
            score = f"{sum(x['score'] for x in assessment['criteria'])}/20" if assessment else "미평가"
            gaps = ([f"{x['criterion']}: {x['judgement']}" for x in assessment["criteria"] if x["score"] < 4]
                    if assessment else [])
            selected = ([c for c in assessment["criteria"] if c["score"] < 4] or assessment["criteria"][:2]) if assessment else []
            reason = [f"{c['criterion']}: {c['observation']} → {c['judgement']} (근거: {c['evidence_ref']})" for c in selected[:2]]
            values.append("\n".join([f"{req['model']} · 추론 {req['reasoning']}",
                                      f"품질 {score} · {_cost(res, req['service'], rates)} · {_time(res)}",
                                      f"상태 {_text(res.get('status'))}", *reason]))
        narrative = comparisons.get((group, case_id))
        values.append(narrative['conclusion'] + '\n' + '\n'.join(
            f"{o['request_id']}: {o['observation']} ({o['evidence_ref']})" for o in narrative['observations'])
            if narrative else "비교 해설 미기록 — 같은 점수라도 답변 구성·제안·실무 영향의 차이는 별도 검토 필요")
        add(compare, values, 190)
    compare.auto_filter.ref = f"A4:{get_column_letter(3 + len(services))}{max(4, compare.max_row)}"
    compare.sheet_properties.pageSetUpPr.fitToPage = True
    compare.page_setup.fitToWidth = 1
    compare.page_setup.orientation = "landscape"

    setup(detail, ["요청 ID", "모델 묶음", "질문", "서비스", "평가 기준", "점수 /4",
                   "답변에서 확인한 사실", "판단·감점 이유", "근거 위치", "평가자", "블라인드"],
          [27, 23, 22, 18, 32, 10, 62, 62, 35, 19, 13])
    detail_ranges = {}
    for item in rows:
        assessment = item["assessment"]
        if assessment is None:
            continue
        first = detail.max_row + 1
        for criterion in assessment["criteria"]:
            row_no = add(detail, [item["request"]["request_id"], item["group"], item["case"]["id"],
                                  SERVICE_NAMES.get(item["request"]["service"], item["request"]["service"]),
                                  criterion["criterion"], criterion["score"], criterion["observation"],
                                  criterion["judgement"], criterion["evidence_ref"],
                                  assessment["evaluator"], "예" if assessment["blind"] else "아니오"], 75)
            if criterion["score"] < 4:
                detail.cell(row_no, 6).fill = PatternFill(
                    "solid", fgColor="FFF0D7" if criterion["score"] == 3 else "FCE4E2")
        detail_ranges[item["request"]["request_id"]] = (first, detail.max_row)
    detail.auto_filter.ref = f"A4:K{max(4, detail.max_row)}"

    setup(ledger, ["요청 ID", "모델 묶음", "질문", "서비스", "표시 모델", "추론", "비교 범위",
                   "상태", "품질 /20", "원시 차감(검증용)", "원시 단위", "환산율(원/단위)",
                   "비용(KRW)", "비용 근거", "시간 구간", "입력 SHA256", "제한·오류"],
          [27, 23, 22, 18, 30, 15, 18, 19, 13, 20, 16, 20, 17, 20, 34, 66, 50])
    totals, costs = {}, {}
    for item in rows:
        req, res = item["request"], item["result"]
        cost = converted_cost(res, req["service"], rates)
        row_no = add(ledger, [req["request_id"], item["group"], item["case"]["id"],
                              SERVICE_NAMES.get(req["service"], req["service"]), req["model"], req["reasoning"],
                              req["comparison_tier"], _text(res.get("status")), None,
                              float(cost["raw"]) if cost else None, res.get("cost_unit"),
                              float(cost["rate"]) if cost else None, None, res.get("cost_status"), _time(res),
                              req["prompt_sha256"], res.get("limitations")], 58)
        if cost:
            ledger.cell(row_no, 13, f"=J{row_no}*L{row_no}")
            ledger.cell(row_no, 13).number_format = '"₩"#,##0.00'
            ledger.cell(row_no, 13).font = Font(name="Arial", bold=True, color=navy)
            costs[row_no] = cost["krw"]
        if req["request_id"] in detail_ranges:
            first, last = detail_ranges[req["request_id"]]
            ledger.cell(row_no, 9, f"=SUM('기준별 평가'!F{first}:F{last})")
            ledger.cell(row_no, 9).font = Font(name="Arial", bold=True, color=navy)
            totals[row_no] = sum(c["score"] for c in item["assessment"]["criteria"])
    ledger.auto_filter.ref = f"A4:Q{max(4, ledger.max_row)}"

    setup(answers, ["요청 ID", "질문 ID", "서비스", "구간", "공통 질문 원문", "답변 원문", "모델", "추론"],
          [27, 23, 18, 10, 95, 115, 28, 15])
    for item in rows:
        req, res = item["request"], item["result"]
        answer = res.get("answer") or "미수집"
        for part, start in enumerate(range(0, len(answer), 1400), 1):
            add(answers, [req["request_id"], item["case"]["id"],
                          SERVICE_NAMES.get(req["service"], req["service"]), part,
                          req["prompt"] if part == 1 else "앞 행과 동일", answer[start:start + 1400],
                          req["model"], req["reasoning"]], 135)
    answers.auto_filter.ref = f"A4:H{max(4, answers.max_row)}"

    setup(ux, ["서비스", "관찰 대상", "실제 관찰", "화면·기록 근거"], [20, 35, 90, 65])
    for event in doc["events"]:
        if event["kind"] == "ux":
            for observation in event["observations"]:
                add(ux, [SERVICE_NAMES.get(event["service"], event["service"]),
                         observation["aspect"], observation["observation"], observation["evidence"]], 72)
    if ux.max_row == 4:
        add(ux, ["전체", "미관측", "UX 관측 기록이 없습니다.", "점수로 추정하지 않음"])
    ux.auto_filter.ref = f"A4:D{ux.max_row}"
    overview = wb.create_sheet("모델별 요약", 0)
    setup(overview, ["모델", "진행 범위", "서비스", "공통 질문 수", "품질 평균 /20",
                     "공통 질문 비용 합계(KRW)", "완료 관측 상한 중앙값", "만점/평가", "답변 차이·평가 근거"],
          [26, 40, 20, 14, 24, 24, 24, 16, 100])
    overview.cell(3, 1, "생성 시점 요약. 원장 변경 시 report 재생성. 파일 생성·탐색 조건은 평균에서 제외. 속도 순위 아님.")
    for group in report_groups(rows, doc):
        for entry in summaries(group, rates):
            add(overview, [group['label'], coverage_text(group), SERVICE_NAMES.get(entry['service'], entry['service']),
                           entry['n'], entry['quality'], entry['cost'], entry['time'], entry['perfect'], entry['evidence']], 145)
        for service, reason in group['excluded'].items():
            add(overview, [group['label'], "설정상 비교 대상 제외 · 현재 메뉴 근거 확인 필요",
                           SERVICE_NAMES.get(service, service), None, None, None, None, None, reason], 50)
    for warning in assessment_warnings(rows):
        add(overview, ["평가 검토", warning], 55)
    if doc.get("report_scope", {}).get("origin") != "init_snapshot":
        add(overview, ["범위 주의", "실행 시작 시 계획 미고정. 현재 설정을 참고한 누락 점검이며 과거 승인 범위를 의미하지 않습니다."], 65)
    overview.freeze_panes = "C5"
    overview.auto_filter.ref = f"A4:I{overview.max_row}"
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)
    _cache_ledger_totals(path, totals, costs, wb.sheetnames.index("실측 원장") + 1)
    with zipfile.ZipFile(path) as archive:
        if archive.testzip() is not None:
            raise ValueError("생성한 XLSX 압축 구조가 손상됨")


def _write_pdf(path, rows, doc, today, name, rates):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak

    font = "BenchmarkKorean"
    candidates = ["/System/Library/Fonts/Supplemental/Arial Unicode.ttf",
                  "/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc", "C:/Windows/Fonts/malgun.ttf"]
    for candidate in candidates:
        if Path(candidate).is_file():
            try:
                pdfmetrics.registerFont(TTFont(font, candidate))
                break
            except Exception:
                continue
    else:
        font = "HYSMyeongJo-Medium"
        pdfmetrics.registerFont(UnicodeCIDFont(font))
    ink, navy = colors.HexColor("#203047"), colors.HexColor("#193B63")
    body = ParagraphStyle("body", fontName=font, fontSize=9, leading=14, textColor=ink, wordWrap="CJK")
    small = ParagraphStyle("small", parent=body, fontSize=8, leading=12)
    heading = ParagraphStyle("heading", parent=body, fontSize=18, leading=26, textColor=navy, spaceAfter=12)
    sub = ParagraphStyle("sub", parent=body, fontSize=12, leading=19, textColor=navy,
                         spaceBefore=15, spaceAfter=7)

    def p(value, style=body):
        return Paragraph(escape(_text(value)).replace("\n", "<br/>"), style)

    def table(matrix, widths):
        result = Table([[p(value, small) for value in row] for row in matrix],
                       colWidths=widths, repeatRows=1, hAlign="LEFT")
        result.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E9F0F8")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ("LINEBELOW", (0, 0), (-1, -1), .35, colors.HexColor("#D9E2EC")),
        ]))
        return result

    rates_label = ", ".join(f"{SERVICE_NAMES.get(service, service)} 1크레딧={krw_text(rate)}"
                            for service, rate in rates.items()) or "적용된 크레딧 환산율 없음"
    story = [p("AI 서비스 비교 보고서", heading),
             p(f"생성일 {today.isoformat()} · 실행 {name} · 결과 {len(rows)}건"),
             Spacer(1, 10), p("읽는 법", sub),
             p("표시 모델·추론·입력 조건의 서비스 경험 비교입니다. 실제 백엔드 동일성은 미검증입니다. "
               "품질은 질문별 5개 기준(각 0~4점)으로 평가하며 비용·속도·UX와 합산하지 않습니다. "
               "시간의 ≤는 완료 관측 상한만 확인했다는 뜻입니다. 비용은 사용자가 제공한 서비스별 원/크레딧 환산율을 적용한 원화입니다. "
               "환산율은 실제 결제액·할인·구독료와 다를 수 있습니다. 미노출·미확정은 0원이 아닙니다."),
             p("적용 환산율: " + rates_label, small)]
    story.append(p("모델별 실행 범위", sub))
    coverage = [["모델", "결과 기록 / 대상", "예정 질문"]]
    for group in report_groups(rows, doc):
        coverage.append([group['label'], coverage_text(group), ", ".join(group['cases'])])
    story.append(table(coverage, [110, 170, 227]))
    if doc.get("report_scope", {}).get("origin") != "init_snapshot":
        story.append(p("시작 시 계획 미고정: 현재 설정을 참고한 누락 점검입니다. 과거 승인 범위나 전체 완료를 뜻하지 않습니다.", small))
    for warning in assessment_warnings(rows):
        story.append(p(warning, small))
    for group in report_groups(rows, doc):
        story.extend([PageBreak(), p(group['label'] + " · 서비스 비교 요약", heading), p(coverage_text(group))])
        story.append(p("같은 질문에서 모든 대상 서비스의 성공·평가·UI 조건 일치가 확인된 교집합만 집계합니다. 파일 생성은 별도 기능 평가이며, 비용 미노출은 무료가 아닙니다. 완료 관측 상한 중앙값은 실제 응답시간 순위가 아닙니다.", small))
        matrix = [["서비스", "공통 질문", "품질 평균", "비용 합계", "관측 상한 중앙값", "만점/평가"]]
        entries = summaries(group, rates)
        for entry in entries:
            matrix.append([SERVICE_NAMES.get(entry['service'], entry['service']), entry['n'], entry['quality'],
                           entry['cost'], entry['time'], entry['perfect']])
        story.append(table(matrix, [83, 51, 110, 80, 103, 80]))
        for entry in entries:
            story.extend([p(SERVICE_NAMES.get(entry['service'], entry['service']) + " · 확인된 차이", sub), p(entry['evidence'])])
        for service, reason in group['excluded'].items():
            story.append(p(f"{SERVICE_NAMES.get(service, service)}: 설정상 비교 대상 제외 — {reason}", small))
    story.extend([PageBreak(), p("질문별 실측 목록", heading)])
    summary = [["질문", "서비스", "표시 모델 · 추론", "상태 · 품질", "비용(KRW)", "시간"]]
    for item in rows:
        req, res, assessment = item["request"], item["result"], item["assessment"]
        score = f"{sum(c['score'] for c in assessment['criteria'])}/20" if assessment else "미평가"
        summary.append([item["case"]["id"], SERVICE_NAMES.get(req["service"], req["service"]),
                        req["model"] + " · " + req["reasoning"], _text(res.get("status")) + " · " + score,
                        _cost(res, req["service"], rates), _time(res)])
    groups = {}
    for item in rows:
        groups.setdefault((item["group"], item["case"]["id"]), {})[item["request"]["service"]] = item
    group_members = {}
    for item in rows:
        group_members.setdefault(item["group"], []).append(item)
    planned_cases = {g['label']: set(g['cases']) for g in report_groups(rows, doc)}
    for group in group_members:
        for case in doc["cases"]:
            if group in planned_cases and case['id'] not in planned_cases[group]:
                continue
            groups.setdefault((group, case["id"]), {})
    all_services = list(SERVICE_NAMES) + list(dict.fromkeys(
        item["request"]["service"] for item in rows if item["request"]["service"] not in SERVICE_NAMES))
    for (group, case_id), members in groups.items():
        for service in all_services:
            if service in members:
                continue
            exclusion = next((r["request"].get("excluded_services", {}).get(service)
                              for r in group_members[group]
                              if r["request"].get("excluded_services", {}).get(service)), None)
            summary.append([case_id, SERVICE_NAMES.get(service, service),
                            group + (" · 모델 없음" if exclusion else " · 미실측"),
                            exclusion or "기록 없음 — 모델 없음으로 추정하지 않음",
                            "해당 없음", "해당 없음"])
    story.append(table(summary, [72, 74, 135, 87, 75, 64]))
    story.extend([Spacer(1, 10),
                  p("각 판단의 실제 관찰·감점 이유는 아래에, 전체 답변 원문은 동봉한 엑셀에 있습니다.", small)])
    comparisons = [e for e in doc['events'] if e['kind'] == 'comparison']
    story.extend([PageBreak(), p("서비스 간 답변 차이와 실무 영향", heading)])
    if not comparisons:
        story.append(p("비교 해설 미기록. 개별 채점표만 있으며 동점 답변의 제안·구성·실무 영향은 평가되지 않았습니다. 상세 채점표의 존재만으로 비교 분석 완료라고 보지 않습니다."))
    for event in comparisons:
        story.extend([p(event['model_group'] + ' · ' + event['case_id'], sub), p(event['conclusion'])])
        story.append(table([["요청", "실제 차이", "근거"]] +
                           [[o['request_id'], o['observation'], o['evidence_ref']] for o in event['observations']],
                           [120, 260, 127]))
    for item in rows:
        req, res, assessment = item["request"], item["result"], item["assessment"]
        story.extend([PageBreak(),
                      p(f"{item['group']} · {item['case']['id']} · {SERVICE_NAMES.get(req['service'], req['service'])}", heading),
                      p(f"{req['model']} / 추론 {req['reasoning']} / {_text(res.get('status'))} / "
                        f"{_cost(res, req['service'], rates)} / {_time(res)}"),
                      p("답변 평가", sub)])
        if assessment:
            story.append(p(f"평가자: {assessment['evaluator']} · 블라인드: {'예' if assessment['blind'] else '아니오'}", small))
            criteria = [["평가 기준", "점수", "확인한 사실과 위치", "판단·감점 이유"]]
            for c in assessment["criteria"]:
                criteria.append([c["criterion"], f"{c['score']}/4",
                                 c["observation"] + "\n근거: " + c["evidence_ref"], c["judgement"]])
            story.append(table(criteria, [111, 42, 176, 178]))
        else:
            story.append(p("성공 답변이 아니므로 품질 점수는 미평가. 실패·비용 상태는 별도로 유지합니다."))
        if res.get("limitations"):
            story.extend([p("제한과 오류", sub), p(res["limitations"])])
        if res.get("status") != "success" or item['case']['id'] == 'file-generation-01':
            story.extend([p("기능 결과·검증 근거", sub), p(res.get("answer")), p(res.get("evidence"), small)])
    story.extend([PageBreak(), p("공통 질문 원문", heading)])
    for case in doc['cases']:
        story.extend([p(case['id'], sub), p("\n".join(case['turns']), small)])
    story.extend([PageBreak(), p("사용 경험과 관찰", heading)])
    ux_events = [event for event in doc["events"] if event["kind"] == "ux"]
    if not ux_events:
        story.append(p("UX 관측 기록이 없어 평가하지 않았습니다."))
    for event in ux_events:
        story.append(p(SERVICE_NAMES.get(event["service"], event["service"]), sub))
        story.append(table([["관찰 대상", "실제 관찰", "근거"]] +
                           [[x["aspect"], x["observation"], x["evidence"]] for x in event["observations"]],
                           [100, 295, 112]))
    notes = [event["text"] for event in doc["events"] if event["kind"] == "note"]
    if notes:
        story.append(p("정정·한계", sub))
        story.extend(p(note) for note in notes)

    def footer(canvas, current):
        canvas.setFont(font, 8)
        canvas.setFillColor(colors.HexColor("#61738A"))
        canvas.drawString(44, 27, f"개인 비교 보고서 · {today.isoformat()} · {name}")
        canvas.drawRightString(A4[0] - 44, 27, str(current.page))

    SimpleDocTemplate(str(path), pagesize=A4, leftMargin=44, rightMargin=44,
                      topMargin=40, bottomMargin=43, title="AI 서비스 비교 보고서",
                      author="").build(story, onFirstPage=footer, onLaterPages=footer)
    if not path.read_bytes().startswith(b"%PDF"):
        raise ValueError("생성한 PDF 파일이 유효하지 않음")


def _write_chart_pdf(path, rows, today, name, rates, doc=None):
    from .charts import write_chart_pdf
    write_chart_pdf(path, rows, today, name, rates, doc=doc)


def report(root, name, *, report_date=None, credit_rates=None, output_dir=None):
    """Export a private dated report, workbook, and chart; never replace prior files."""
    today = date.fromisoformat(report_date) if report_date else date.today()
    rates = normalize_rates(credit_rates)
    doc = read_session(root, name)
    if "report_scope" not in doc and (Path(root) / "config/benchmark.json").is_file():
        from .planning import read_json
        doc["report_scope"] = {"models": read_json(root, "config/benchmark.json")["models"],
                               "origin": "current_config_not_original_plan"}
    rows = _rows(doc)
    if not rows:
        raise ValueError("기록된 요청이 없어 최종 보고서를 만들 수 없음")
    # Validate every known charge before creating any final files.
    for item in rows:
        converted_cost(item["result"], item["request"]["service"], rates)
    folder = (Path(output_dir).expanduser() if output_dir is not None
              else Path.home() / "AI-Benchmark-Reports" / name).resolve()
    folder.mkdir(parents=True, exist_ok=True)
    stem = f"AI-service-comparison_{today.isoformat()}_{name}"
    pdf, xlsx, chart = (folder / f"{stem}.pdf", folder / f"{stem}.xlsx",
                        folder / f"{stem}_chart.pdf")
    if pdf.exists() or xlsx.exists() or chart.exists():
        raise FileExistsError("최종 파일이 이미 있습니다. 기존 파일을 확인하고 새 실행 ID를 사용하세요")
    try:
        with tempfile.TemporaryDirectory(dir=folder) as temporary:
            draft_pdf, draft_xlsx, draft_chart = (Path(temporary) / "draft.pdf",
                                                  Path(temporary) / "draft.xlsx",
                                                  Path(temporary) / "chart.pdf")
            _write_xlsx(draft_xlsx, rows, doc, today, name, rates)
            _write_pdf(draft_pdf, rows, doc, today, name, rates)
            _write_chart_pdf(draft_chart, rows, today, name, rates, doc)
            os.replace(draft_xlsx, xlsx)
            os.replace(draft_pdf, pdf)
            os.replace(draft_chart, chart)
    except ImportError as exc:
        raise ValueError("PDF·XLSX 출력에는 requirements-report.txt 설치가 필요합니다") from exc
    return {"report": str(pdf), "workbook": str(xlsx), "chart": str(chart), "records": len(rows),
            "private": True, "external_upload": False}
