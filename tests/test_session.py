import copy
import hashlib
import importlib.util
import json
import tempfile
import unittest
from pathlib import Path
from benchmark.session import initialize, append_event, status, session_dir
from benchmark.reporting import report
from benchmark.costs import converted_cost, normalize_rates


class SessionTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        (self.root / 'cases').mkdir()
        (self.root / 'cases/core.json').write_text(json.dumps({'cases':[
            {'id':'summary-01','turns':['<script>hi</script>'],'fixture_ids':[],
             'rubric':['정확성','근거']},
            {'id':'pending-01','turns':['pending'],'fixture_ids':[],
             'rubric':['정확성','근거'] }]}))
        initialize(self.root, 'test-run')
        self.event = dict(kind='reserve',request_id='mono-summary',service='monogpt',
            model='common-model',reasoning='low',case_id='summary-01',
            ui_evidence='UI selected low',authorization_scope='limited pilot acknowledged',
            prompt='<script>hi</script>',prompt_sha256=hashlib.sha256(b'<script>hi</script>').hexdigest(),
            comparison_tier='ui_matched',fixtures={},feature_settings={'search':'off'},
            request_limit=2,cost_mode='request_limited_unknown_price',unknown_price_acknowledged=True)
    def success(self):
        append_event(self.root,'test-run',self.event)
        append_event(self.root,'test-run',dict(kind='result',request_id='mono-summary',status='success',
            answer='=HYPERLINK("bad") <script>alert(1)</script>',cost_status='unavailable',cost=None))
    def assessment(self):
        return dict(kind='assessment',request_id='mono-summary',evaluator='test reviewer',blind=False,
                    criteria=[dict(criterion='정확성',score=3,observation='핵심 내용을 확인함',
                                   judgement='한 조건이 누락됨',evidence_ref='답변 첫 문장'),
                              dict(criterion='근거',score=4,observation='근거 위치가 명시됨',
                                   judgement='요구 근거를 충족함',evidence_ref='답변 둘째 문장')])
    def test_assessment_requires_success_and_evidence(self):
        append_event(self.root,'test-run',self.event)
        with self.assertRaises(ValueError):
            append_event(self.root,'test-run',self.assessment())
        append_event(self.root,'test-run',dict(kind='result',request_id='mono-summary',status='success',
            answer='answer',cost_status='unavailable',cost=None))
        bad=self.assessment()
        bad['criteria'][0]['evidence_ref']=''
        with self.assertRaises(ValueError):
            append_event(self.root,'test-run',bad)
        append_event(self.root,'test-run',self.assessment())
        with self.assertRaises(ValueError):
            append_event(self.root,'test-run',self.assessment())
    def test_ux_requires_observation_and_evidence(self):
        with self.assertRaises(ValueError):
            append_event(self.root,'test-run',dict(kind='ux',service='monogpt',
                observations=[dict(aspect='가독성',observation='좋음',evidence='')]))
        append_event(self.root,'test-run',dict(kind='ux',service='monogpt',
            observations=[dict(aspect='가독성',observation='표의 제목과 숫자가 구분됨',evidence='화면 캡처 1')]))
    def test_report_blocks_unassessed_success(self):
        self.success()
        with self.assertRaises(ValueError):
            report(self.root,'test-run',output_dir=self.root/'exports',report_date='2026-10-15')
    @unittest.skipUnless(importlib.util.find_spec('openpyxl') and importlib.util.find_spec('reportlab'),
                         'PDF/XLSX optional dependencies not installed')
    def test_dated_pdf_xlsx_and_literal_answer(self):
        from openpyxl import load_workbook
        self.event['excluded_services']={'chatgpt':'synthetic not offered'}
        self.success()
        append_event(self.root,'test-run',self.assessment())
        output=report(self.root,'test-run',output_dir=self.root/'exports',report_date='2026-10-15')
        pdf,xlsx,chart=Path(output['report']),Path(output['workbook']),Path(output['chart'])
        self.assertEqual(pdf.parent,xlsx.parent)
        self.assertEqual(pdf.parent,(self.root/'exports').resolve())
        self.assertEqual(sorted(p.name for p in pdf.parent.iterdir()),sorted([pdf.name,xlsx.name,chart.name]))
        self.assertIn('2026-10-15',pdf.name)
        self.assertTrue(pdf.read_bytes().startswith(b'%PDF'))
        self.assertTrue(chart.read_bytes().startswith(b'%PDF'))
        book=load_workbook(xlsx)
        self.assertEqual(book['기준별 평가']['F5'].value,3)
        self.assertIn('모델 없음',book['서비스 비교']['E5'].value)
        self.assertEqual(book['실측 원장']['I5'].value,"=SUM('기준별 평가'!F5:F6)")
        self.assertEqual(load_workbook(xlsx,data_only=True)['실측 원장']['I5'].value,7)
        self.assertEqual(book['답변 원문']['F5'].data_type,'s')
        self.assertTrue(book['답변 원문']['F5'].value.startswith('=HYPERLINK'))
        with self.assertRaises(FileExistsError):
            report(self.root,'test-run',output_dir=self.root/'exports',report_date='2026-10-15')
    def test_credit_conversion_requires_user_rate_and_preserves_unknown(self):
        self.assertIsNone(converted_cost({'cost':None,'cost_unit':'credits'},'monogpt',{}))
        with self.assertRaisesRegex(ValueError, '환산율'):
            converted_cost({'cost':'58','cost_unit':'credits'},'monogpt',{})
        self.assertEqual(converted_cost({'cost':'58','cost_unit':'credits'},'monogpt',
                                        normalize_rates({'monogpt':'1'}))['krw'],58)
        self.assertEqual(converted_cost({'cost':'43','cost_unit':'크레딧'},'timely',
                                        normalize_rates({'timely':'2'}))['krw'],86)
        self.assertEqual(converted_cost({'cost':'10','cost_unit':'KRW'},'chatgpt',{})['krw'],10)
        with self.assertRaisesRegex(ValueError, '0보다 큰'):
            normalize_rates({'timely':'0'})
        with self.assertRaisesRegex(ValueError, '알 수 없는'):
            converted_cost({'cost':'10','cost_unit':'tokens'},'timely',{'timely':2})
    @unittest.skipUnless(importlib.util.find_spec('openpyxl') and importlib.util.find_spec('reportlab'),
                         'PDF/XLSX optional dependencies not installed')
    def test_report_converts_credit_to_krw_in_all_outputs(self):
        from openpyxl import load_workbook
        append_event(self.root,'test-run',self.event)
        append_event(self.root,'test-run',dict(kind='result',request_id='mono-summary',status='success',
            answer='검증 가능한 답변',cost_status='ledger_verified',cost='58',cost_unit='credits'))
        append_event(self.root,'test-run',self.assessment())
        timely=dict(self.event,request_id='timely-summary',service='timely')
        append_event(self.root,'test-run',timely)
        append_event(self.root,'test-run',dict(kind='result',request_id='timely-summary',status='success',
            answer='타임리 답변',cost_status='balance_observed',cost='43',cost_unit='크레딧'))
        append_event(self.root,'test-run',dict(self.assessment(),request_id='timely-summary'))
        with self.assertRaisesRegex(ValueError,'환산율'):
            report(self.root,'test-run',output_dir=self.root/'exports',report_date='2026-10-15')
        with self.assertRaisesRegex(ValueError,'timely'):
            report(self.root,'test-run',output_dir=self.root/'exports',report_date='2026-10-15',credit_rates={'monogpt':'1'})
        output=report(self.root,'test-run',output_dir=self.root/'exports',report_date='2026-10-15',
                      credit_rates={'monogpt':'1','timely':'2'})
        book=load_workbook(output['workbook'])
        ledger=book['실측 원장']
        self.assertEqual(ledger['M5'].value,'=J5*L5')
        self.assertEqual(ledger['L5'].value,1)
        self.assertEqual(load_workbook(output['workbook'],data_only=True)['실측 원장']['M5'].value,58)
        self.assertEqual(load_workbook(output['workbook'],data_only=True)['실측 원장']['M6'].value,86)
        self.assertIn('58원',book['서비스 비교']['C5'].value)
        self.assertIn('86원',book['서비스 비교']['D5'].value)
        self.assertTrue(Path(output['report']).read_bytes().startswith(b'%PDF'))
    @unittest.skipUnless(importlib.util.find_spec('openpyxl') and importlib.util.find_spec('reportlab'),
                         'PDF/XLSX optional dependencies not installed')
    def test_default_report_folder_is_visible_home_folder(self):
        from unittest.mock import patch
        self.success()
        append_event(self.root,'test-run',self.assessment())
        with patch('benchmark.reporting.Path.home',return_value=self.root):
            output=report(self.root,'test-run',report_date='2026-10-15')
        self.assertEqual(Path(output['report']).parent,(self.root/'AI-Benchmark-Reports'/'test-run').resolve())
    def test_duplicate_and_unresolved_block(self):
        append_event(self.root,'test-run',self.event)
        with self.assertRaises(ValueError): append_event(self.root,'test-run',self.event)
        other=dict(self.event,request_id='other')
        with self.assertRaises(ValueError): append_event(self.root,'test-run',other)
    def test_prompt_hash(self):
        with self.assertRaises(ValueError): append_event(self.root,'test-run',dict(self.event,prompt_sha256='bad'))
    def test_unknown_reasoning(self):
        with self.assertRaises(ValueError): append_event(self.root,'test-run',dict(self.event,reasoning='not_exposed'))
    def test_condition_mismatch(self):
        append_event(self.root,'test-run',self.event)
        with self.assertRaises(ValueError): append_event(self.root,'test-run',dict(self.event,request_id='timely-summary',service='timely',reasoning='high'))
    def test_cost_unavailable_not_zero(self):
        append_event(self.root,'test-run',self.event)
        with self.assertRaises(ValueError): append_event(self.root,'test-run',dict(kind='result',request_id='mono-summary',status='timeout',cost_status='unavailable',cost=0))
    def test_no_path_escape_or_overwrite(self):
        with self.assertRaises(ValueError): session_dir(self.root,'../outside')
        with self.assertRaises(FileExistsError): initialize(self.root,'test-run')
    def test_limit(self):
        append_event(self.root,'test-run',dict(self.event,request_limit=1))
        with self.assertRaises(ValueError): append_event(self.root,'test-run',dict(self.event,request_id='timely-summary',service='timely'))
    def test_limit_accepts_40_but_not_41(self):
        append_event(self.root,'test-run',dict(self.event,request_limit=40))
        other=dict(self.event,request_id='timely-summary',service='timely',request_limit=41)
        with self.assertRaisesRegex(ValueError,'1~40'):
            append_event(self.root,'test-run',other)
    def test_result_requires_reservation(self):
        with self.assertRaises(ValueError): append_event(self.root,'test-run',dict(kind='result',request_id='absent',status='timeout',cost_status='unavailable'))
